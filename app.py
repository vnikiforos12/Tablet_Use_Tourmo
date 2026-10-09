import io
import json
import re
import time
import urllib.request
import chardet
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

# ==============================================================================
# 1. ΡΥΘΜΙΣΗ ΣΕΛΙΔΑΣ & ΕΤΑΙΡΙΚΟ ΣΤΥΛ ΟΜΙΛΟΥ ΗΡΑΚΛΗΣ
# ==============================================================================
st.set_page_config(
    page_title="Όμιλος ΗΡΑΚΛΗΣ | Logistics & Telematics Reconciler",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
        .main { background-color: #f8fafc; }
        .heracles-header {
            background: linear-gradient(135deg, #002B49 0%, #004b7a 60%, #0072CE 100%);
            padding: 22px 28px;
            border-radius: 12px;
            color: white;
            margin-bottom: 20px;
            border-left: 8px solid #00A3E0;
        }
        .heracles-header h1 { color: #ffffff; font-size: 26px; font-weight: 800; margin: 0; }
        .heracles-header p { color: #dbeafe; font-size: 14px; margin-top: 5px; margin-bottom: 0; }
        .stButton>button {
            background: linear-gradient(135deg, #002B49 0%, #0072CE 100%) !important;
            color: white !important;
            font-weight: 700 !important;
            font-size: 16px !important;
            padding: 12px 28px !important;
            border-radius: 8px !important;
            border: none !important;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

# Keep-Alive Heartbeat
components.html(
    """<script>setInterval(function(){window.dispatchEvent(new Event('resize'));fetch(window.location.href,{mode:'no-cors'}).catch(()=>{});},40000);</script>""",
    height=0,
)

st.markdown(
    """
    <div class="heracles-header">
        <h1>🏛️ ΟΜΙΛΟΣ ΗΡΑΚΛΗΣ | Logistics & Telematics Reconciler</h1>
        <p>Έκδοση 4.1 (Χρωματισμός Ποσοστών: 🟢 ≥90% Πράσινο, 🟡 <30% Κίτρινο, 🔴 0% Κόκκινο)</p>
    </div>
    """,
    unsafe_allow_html=True,
)

# ==============================================================================
# 2. ΧΑΡΤΟΓΡΑΦΗΣΗ ΑΠΟΣΤΑΣΕΩΝ & ΠΛΟΙΩΝ
# ==============================================================================
SHIPPING_COORDS = {
    "Μηλάκι οδικές φορτώσεις": (38.379005, 24.065336),
    "ΚΔ ΔΡΑΠΕΤΣΩΝΑΣ": (37.947882, 23.621671),
    "ΒΟΛΟΣ_ΟΔΙΚΕΣ_ΦΟΡΤΩΣΕΙΣ": (39.351592, 22.984101),
    "Thessaloniki Terminal": (40.645465, 22.898439),
    "ΚΔ ΗΡΑΚΛΕΙΟΥ": (35.348496, 25.044322),
    "Kavala Terminal": (40.926580, 24.389990),
    "Igoumenitsa Terminal": (39.486899, 20.248540),
    "Thessaloniki Terminal-Exports": (40.645465, 22.898439),
    "ΚΔ ΡΙΟΥ": (38.311128, 21.793005),
}

SHIPPING_TO_PARENT = {
    "Μηλάκι οδικές φορτώσεις": "Milaki",
    "ΚΔ ΔΡΑΠΕΤΣΩΝΑΣ": "Drapetsona",
    "ΒΟΛΟΣ_ΟΔΙΚΕΣ_ΦΟΡΤΩΣΕΙΣ": "Volos",
    "Thessaloniki Terminal": "Thessaloniki",
    "Thessaloniki Terminal-Exports": "Thessaloniki",
    "ΚΔ ΗΡΑΚΛΕΙΟΥ": "Heraklion",
    "Kavala Terminal": "Kavala",
    "Igoumenitsa Terminal": "Igoumenitsa",
    "ΚΔ ΡΙΟΥ": "Rio",
}

FERRY_DESTINATIONS = [
    {
        "names": ["ΜΥΤΙΛΗΝΗ", "MYTILINI", "KALONH", "ΚΑΛΛΟΝΗ", "ΛΕΣΒΟΣ"],
        "mainland_port": (37.945, 23.635),
        "island_port": (39.104, 26.557),
    },
    {
        "names": ["ΧΙΟΣ", "ΧΑΛΚΕΙΟΣ", "ΘΥΜΙΑΝΑ"],
        "mainland_port": (37.945, 23.635),
        "island_port": (38.371, 26.139),
    },
    {
        "names": ["ΚΥΘΗΡΑ", "KYTHIRA"],
        "mainland_port": (36.511, 23.059),
        "island_port": (36.262, 23.003),
    },
    {
        "names": ["KEA", "ΚΕΑ", "ΤΖΙΑ"],
        "mainland_port": (37.714, 24.058),
        "island_port": (37.658, 24.310),
    },
    {
        "names": ["ΑΝΔΡΟΣ", "ΑΝΔΡΟΥ", "ANDROS"],
        "mainland_port": (38.022, 24.008),
        "island_port": (37.883, 24.736),
    },
    {
        "names": ["ΣΥΡΟΣ", "SYROS"],
        "mainland_port": (37.945, 23.635),
        "island_port": (37.440, 24.943),
    },
    {
        "names": ["ΠΑΡΟΣ", "ΠΑΡΟΥ", "ΑΓΚΑΙΡΙΑ", "PAROS"],
        "mainland_port": (37.945, 23.635),
        "island_port": (37.085, 25.150),
    },
    {
        "names": ["ΝΑΞΟΣ", "NAKSOS", "NAXOS"],
        "mainland_port": (37.945, 23.635),
        "island_port": (37.107, 25.372),
    },
    {
        "names": ["ΜΗΛΟΣ", "MILOS", "ΑΔΑΜΑΣ"],
        "mainland_port": (37.945, 23.635),
        "island_port": (36.724, 24.446),
    },
    {
        "names": ["ΣΙΦΝΟΣ", "ΣΙΦΝΟΥ", "SIFNOS"],
        "mainland_port": (37.945, 23.635),
        "island_port": (36.989, 24.675),
    },
    {
        "names": ["ΘΗΡΑΣ", "ΘΗΡΑ", "ΣΑΝΤΟΡΙΝΗ", "SANTORINI", "ΚΑΜΑΡΙ"],
        "mainland_port": (37.945, 23.635),
        "island_port": (36.386, 25.431),
    },
    {
        "names": [
            "ΚΕΡΚΥΡΑ",
            "ΑΧΑΡΑΒΗ",
            "ΒΑΡΥΠΑΤΑΔΕΣ",
            "ΔΑΝΙΛΙΑ",
            "ΜΕΣΟΓΓΗ",
            "BELONADES",
            "ΧΩΡΟΕΠΙΣΚΟΠΟΙ",
            "CORFU",
        ],
        "mainland_port": (39.497, 20.258),
        "island_port": (39.627, 19.907),
    },
    {
        "names": ["ΘΑΣΟΣ", "ΘΑΣΟ", "THASOS"],
        "mainland_port": (40.856, 24.700),
        "island_port": (40.780, 24.710),
    },
    {
        "names": ["ΣΚΙΑΘΟΣ", "SKIATHOS"],
        "mainland_port": (39.356, 22.946),
        "island_port": (39.163, 23.491),
    },
    {
        "names": ["ΣΑΛΑΜΙΝΑ", "ΣΑΛΑΜΙΝΑΣ", "ΜΠΑΤΣΙ"],
        "mainland_port": (37.962, 23.570),
        "island_port": (37.968, 23.535),
    },
    {
        "names": [
            "ΚΑΡΥΕΣ",
            "ΑΓΙΟ ΟΡΟΣ",
            "ΑΓΙΟΥ ΟΡΟΥΣ",
            "ΒΑΤΟΠΕΔΙ",
            "ΞΗΡΟΠΟΤΑΜΟΣ",
        ],
        "mainland_port": (40.326, 23.982),
        "island_port": (40.214, 24.218),
    },
    {
        "names": ["ΙΣΤΙΑΙΑ", "ΑΙΔΗΨΟΣ", "ISTIAIA"],
        "mainland_port": (38.950, 22.964),
        "island_port": (38.928, 23.033),
        "only_from": ["ΒΟΛΟΣ"],
    },
    {
        "names": [
            "ΧΑΝΙΑ",
            "ΧΑΝΙΩΝ",
            "ΗΡΑΚΛΕΙΟ",
            "ΡΕΘΥΜΝΟ",
            "ΑΓΙΟΣ ΝΙΚΟΛΑΟΣ",
            "ΝΕΑΠΟΛΗ",
            "ΙΕΡΑΠΕΤΡΑ",
            "ΜΟΙΡΕΣ",
            "ΚΡΗΤΗ",
        ],
        "mainland_port": (37.945, 23.635),
        "island_port": (35.344, 25.148),
        "exclude_from": ["ΗΡΑΚΛΕΙΟΥ"],
    },
]

routing_cache = {}


def get_osrm_distance(p1, p2):
    lat1, lon1 = p1
    lat2, lon2 = p2
    key = (round(lat1, 4), round(lon1, 4), round(lat2, 4), round(lon2, 4))
    if key in routing_cache:
        return routing_cache[key]
    try:
        url = f"http://router.project-osrm.org/route/v1/driving/{lon1},{lat1};{lon2},{lat2}?overview=false"
        req = urllib.request.Request(
            url, headers={"User-Agent": "HeraclesApp/4.1"}
        )
        with urllib.request.urlopen(req, timeout=10) as response:
            data = json.loads(response.read().decode())
            if data.get("code") == "Ok" and data.get("routes"):
                km = round(data["routes"][0]["distance"] / 1000.0, 2)
                routing_cache[key] = km
                time.sleep(0.1)
                return km
    except Exception:
        pass
    return 0.0


def read_csv_smart(raw_bytes):
    detected = chardet.detect(raw_bytes)
    enc = detected["encoding"] if detected["encoding"] else "utf-8"
    for encoding_type in [enc, "utf-8-sig", "utf-8", "windows-1253"]:
        try:
            return pd.read_csv(
                io.BytesIO(raw_bytes),
                sep=None,
                engine="python",
                encoding=encoding_type,
                dtype=str,
            )
        except Exception:
            continue
    return None


col1, col2 = st.columns(2)
with col1:
    file_zonda = st.file_uploader(
        "1. Αρχείο Zonda CSV (Παραγγελίες)", type=["csv"], key="zonda"
    )
with col2:
    file_tourmo = st.file_uploader(
        "2. Αρχείο TOURMO CSV (Τηλεματική)", type=["csv"], key="tourmo"
    )

if file_zonda and file_tourmo:
    if st.button("🚀 Εκτέλεση Επεξεργασίας & Δημιουργία Αναφοράς"):
        with st.spinner("Επεξεργασία, υπολογισμός αποστάσεων και ενοποίηση..."):
            df_zonda = read_csv_smart(file_zonda.getvalue())
            df_tourmo = read_csv_smart(file_tourmo.getvalue())

            # --- A. ZONDA PROCESSING ---
            cols_to_drop_z = [
                "DELIVERY_FROM_DAT",
                "DELIVERY_TO_DAT",
                "ORDER_NUMBER_FROM_SEQ_USAGE",
                "RETURNED_QUANTITY",
                "RETURNED_UOM",
                "DELIVERED_UOM",
            ]
            found_drop_z = [
                c
                for c in df_zonda.columns
                if any(
                    c.strip("#").strip().lower() == t.lower()
                    for t in cols_to_drop_z
                )
            ]
            if found_drop_z:
                df_zonda.drop(columns=found_drop_z, inplace=True)

            def find_z_col(target):
                for c in df_zonda.columns:
                    if target.lower() in c.lower():
                        return c
                return None

            z_order = find_z_col("ORDER_NUMBER")
            z_status = find_z_col("STATUS")
            z_shipping = find_z_col("SHIPPINGPOINT")
            z_city = find_z_col("CITY")
            z_address = find_z_col("SHIPTO_ADDRESS")
            z_dist = find_z_col("DISTANCE")
            z_qty = find_z_col("DELIVERED_QUANTITY")
            z_sap = find_z_col("COMM_CARR")
            z_carrier = None
            for c in df_zonda.columns:
                if "carrier" in c.lower() and "sap" not in c.lower():
                    z_carrier = c
                    break
            z_vehicle = find_z_col("VEHICLE")
            z_shipto_lat = find_z_col("SHIPTO_LATITUDE")
            z_shipto_lon = find_z_col("SHIPTO_LONGITUDE")

            if z_order:
                cl_order = (
                    df_zonda[z_order]
                    .astype(str)
                    .str.replace("(10)", "", regex=False)
                    .str.strip()
                )
                df_zonda[z_order] = pd.to_numeric(
                    cl_order, errors="coerce"
                ).astype("Int64")

            if z_status:
                df_zonda = df_zonda[
                    ~df_zonda[z_status]
                    .astype(str)
                    .str.strip()
                    .str.lower()
                    .eq("deleted")
                ].copy()

            if z_qty:
                num_qty = pd.to_numeric(
                    df_zonda[z_qty].str.replace(",", "."), errors="coerce"
                )
                df_zonda = df_zonda[num_qty != 0].copy()
                df_zonda[z_qty] = pd.to_numeric(
                    df_zonda[z_qty].str.replace(",", "."), errors="coerce"
                ).round(2)

            if z_vehicle:
                greek_to_latin = {
                    "Α": "A",
                    "Β": "B",
                    "Ε": "E",
                    "Ζ": "Z",
                    "Η": "H",
                    "Ι": "I",
                    "Κ": "K",
                    "Μ": "M",
                    "Ν": "N",
                    "Ο": "O",
                    "Ρ": "P",
                    "Τ": "T",
                    "Υ": "Y",
                    "Χ": "X",
                }

                def clean_plate_fn(val):
                    if pd.isna(val) or str(val).strip() == "":
                        return ""
                    v_str = str(val).strip()
                    pattern = r"([A-Za-zΑ-Ωα-ω]{3}\s*\d{4}|[A-Za-z]{1,2}\s*\d{2,3}\s*[A-Za-z]{3})"
                    m = re.search(pattern, v_str)
                    p = (
                        m.group(1).replace(" ", "").upper()
                        if m
                        else v_str.replace(" ", "").upper()
                    )
                    return "".join(
                        greek_to_latin.get(char, char) for char in p
                    )

                df_zonda[z_vehicle] = df_zonda[z_vehicle].apply(clean_plate_fn)

            if z_sap:
                df_zonda[z_sap] = pd.to_numeric(
                    df_zonda[z_sap], errors="coerce"
                ).astype("Int64")

            if (
                z_sap
                and z_carrier
                and z_sap in df_zonda.columns
                and z_carrier in df_zonda.columns
            ):
                cols = [c for c in df_zonda.columns if c != z_sap]
                c_idx = cols.index(z_carrier)
                cols.insert(c_idx, z_sap)
                df_zonda = df_zonda[cols]

            alerts_zonda = []
            is_dist_changed_zonda = []
            if z_dist:
                final_dist_zonda = []
                for idx, row in df_zonda.iterrows():
                    order_id = row.get(z_order, f"Γραμμή {idx}")
                    city_name = str(row.get(z_city, "")).strip()
                    addr_name = str(row.get(z_address, "")).strip()
                    sp_name = str(row.get(z_shipping, "")).strip()

                    v_raw = str(row[z_dist] or "").strip()
                    c_dist_str = (
                        v_raw.replace("km", "")
                        .replace("KM", "")
                        .strip()
                        .replace(",", ".")
                    )
                    try:
                        orig_d = float(c_dist_str)
                    except ValueError:
                        orig_d = 0.0

                    try:
                        s_lat = float(
                            str(row.get(z_shipto_lat, 0)).replace(",", ".")
                        )
                        s_lon = float(
                            str(row.get(z_shipto_lon, 0)).replace(",", ".")
                        )
                    except ValueError:
                        s_lat, s_lon = 0.0, 0.0

                    comb_loc = f"{city_name} {addr_name}".upper()
                    loc_tokens = set(
                        re.findall(r"[\wΑ-Ωα-ωίϊΐόάέύϋΰήώ]+", comb_loc)
                    )

                    matched_ferry = None
                    for ferry in FERRY_DESTINATIONS:
                        if "exclude_from" in ferry and any(
                            ex.upper() in sp_name.upper()
                            for ex in ferry["exclude_from"]
                        ):
                            continue
                        if "only_from" in ferry and not any(
                            of.upper() in sp_name.upper()
                            for of in ferry["only_from"]
                        ):
                            continue
                        if any(
                            target.upper() in loc_tokens for target in ferry["names"]
                        ):
                            matched_ferry = ferry
                            break

                    if s_lat == 0.0 or s_lon == 0.0:
                        final_dist_zonda.append(0.0)
                        is_dist_changed_zonda.append(False)
                        alerts_zonda.append(
                            f"Παραγγελία {order_id} ({city_name}): Συντεταγμένες 0.0 (Απόσταση 0)"
                        )
                        continue

                    orig_coords = None
                    for sp_k, coords in SHIPPING_COORDS.items():
                        if sp_k.lower() in sp_name.lower():
                            orig_coords = coords
                            break

                    if matched_ferry and orig_coords:
                        isl_port = matched_ferry["island_port"]
                        if (
                            "ΧΑΝΙΑ" in loc_tokens or "CHANIA" in loc_tokens
                        ) and "only_from" not in matched_ferry:
                            isl_port = (35.490, 24.075)

                        d1 = get_osrm_distance(
                            orig_coords, matched_ferry["mainland_port"]
                        )
                        d2 = get_osrm_distance(isl_port, (s_lat, s_lon))
                        tot_km = round((d1 + d2) * 2, 2)
                        final_dist_zonda.append(tot_km)
                        changed = abs(tot_km - round(orig_d * 2, 2)) > 0.5
                        is_dist_changed_zonda.append(changed)
                    elif orig_d == 0.0 and orig_coords:
                        r_km = get_osrm_distance(orig_coords, (s_lat, s_lon))
                        tot_km = round(r_km * 2, 2)
                        final_dist_zonda.append(tot_km)
                        is_dist_changed_zonda.append(tot_km > 0)
                    else:
                        final_dist_zonda.append(round(orig_d * 2, 2))
                        is_dist_changed_zonda.append(False)

                df_zonda[z_dist] = final_dist_zonda

            # --- B. TOURMO PROCESSING ---
            df_tourmo.columns = [c.strip() for c in df_tourmo.columns]

            def find_t_exact(target_name):
                for c in df_tourmo.columns:
                    if c.strip().lower() == target_name.strip().lower():
                        return c
                return None

            def find_t_contains(target_substring):
                for c in df_tourmo.columns:
                    if target_substring.lower() in c.lower():
                        return c
                return None

            t_id = find_t_exact("ID")
            t_name = find_t_exact("Όνομα")
            t_surname = find_t_exact("Επώνυμο")
            t_ext_id = find_t_contains("Εξωτερικό αναγνωριστικό")
            t_phone = find_t_contains("Τηλέφωνο")
            t_dist = find_t_contains("Απόσταση")
            t_unit = find_t_contains("Μονάδα Μέτρησης")

            if t_surname:
                df_tourmo = df_tourmo[
                    ~df_tourmo[t_surname]
                    .astype(str)
                    .str.lower()
                    .str.contains("test")
                ].copy()

            if t_name and t_surname:
                c_name = df_tourmo[t_name].astype(str).str.strip()
                c_sur = df_tourmo[t_surname].astype(str).str.strip()
                comb_v = (
                    (c_name + c_sur)
                    .str.replace(" ", "", regex=False)
                    .str.replace("g", "", case=False, regex=False)
                )

                n_idx = df_tourmo.columns.get_loc(t_name)
                df_tourmo.insert(n_idx, "Όχημα", comb_v)
                df_tourmo.drop(columns=[t_name, t_surname], inplace=True)

            t_ext_id = find_t_contains("Εξωτερικό αναγνωριστικό")
            if t_ext_id:

                def ext_mid_fn(val):
                    if pd.isna(val) or str(val).strip() == "":
                        return ""
                    parts = str(val).strip().split("-")
                    if len(parts) >= 3:
                        return parts[2]
                    m = re.search(r"\b\d{7}\b", str(val))
                    return m.group(0) if m else str(val).strip()

                df_tourmo[t_ext_id] = df_tourmo[t_ext_id].apply(ext_mid_fn)
                df_tourmo[t_ext_id] = pd.to_numeric(
                    df_tourmo[t_ext_id], errors="coerce"
                ).astype("Int64")

            if t_phone and t_phone in df_tourmo.columns:
                df_tourmo.drop(columns=[t_phone], inplace=True)

            t_unit = find_t_contains("Μονάδα Μέτρησης")
            if t_unit and t_unit in df_tourmo.columns:
                u_idx = df_tourmo.columns.get_loc(t_unit)
                df_tourmo = df_tourmo.iloc[:, :u_idx].copy()

            t_dist = find_t_contains("Απόσταση")
            if t_dist and t_dist in df_tourmo.columns:
                df_tourmo[t_dist] = pd.to_numeric(
                    df_tourmo[t_dist].astype(str).str.replace(",", "."),
                    errors="coerce",
                ).round(2)

            if t_id and t_id in df_tourmo.columns:
                df_tourmo[t_id] = pd.to_numeric(
                    df_tourmo[t_id], errors="coerce"
                ).astype("Int64")

            t_parent_group = find_t_contains("Γονική Ομάδα")
            t_group = find_t_exact(
                "Ομάδα"
            )  # Ακριβές όνομα για να μην επιλέγει Εβδομάδα!
            t_veh_col = find_t_exact("Όχημα")

            # 1. Tourmo Pivots
            t_piv1 = (
                df_tourmo.groupby(
                    [t_parent_group, t_group, t_veh_col], as_index=False
                )[t_dist]
                .sum()
                .rename(columns={t_dist: "Tourmo_Km"})
            )
            t_piv1["Tourmo_Km"] = t_piv1["Tourmo_Km"].round(0)

            t_piv2 = (
                df_tourmo.groupby(
                    [t_parent_group, t_ext_id, t_group], as_index=False
                )[t_dist]
                .sum()
                .rename(columns={t_dist: "Tourmo_Km"})
            )
            t_piv2["Tourmo_Km"] = t_piv2["Tourmo_Km"].round(0)

            # --- C. PIVOTS ZONDA (ΜΕ ΓΟΝΙΚΗ ΟΜΑΔΑ ΤΕΡΜΑ ΔΕΞΙΑ) ---
            z_piv1 = (
                df_zonda.groupby([z_carrier, z_vehicle], as_index=False)[z_dist]
                .sum()
                .rename(columns={z_dist: "Zonda_Km"})
            )
            z_piv1["Zonda_Km"] = z_piv1["Zonda_Km"].round(0)
            tourmo_parent_by_veh = (
                df_tourmo.groupby(t_veh_col)[t_parent_group].first().to_dict()
            )
            z_piv1["Γονική Ομάδα"] = (
                z_piv1[z_vehicle].map(tourmo_parent_by_veh).fillna("-")
            )

            z_piv2 = (
                df_zonda.groupby([z_sap, z_carrier], as_index=False)[z_dist]
                .sum()
                .rename(columns={z_dist: "Zonda_Km"})
            )
            z_piv2["Zonda_Km"] = z_piv2["Zonda_Km"].round(0)
            tourmo_parent_by_sap = (
                df_tourmo.groupby(t_ext_id)[t_parent_group].first().to_dict()
            )
            z_piv2["Γονική Ομάδα"] = (
                z_piv2[z_sap].map(tourmo_parent_by_sap).fillna("-")
            )

            # --- D. ΣΥΓΚΡΙΤΙΚΕΣ ΚΑΡΤΕΛΕΣ TABLET USAGE (FULL OUTER JOIN) ---
            def calc_usage_pct(t_val, z_val):
                if z_val == 0:
                    return 1.0  # 100% αν Zonda=0 (είτε Tourmo=0 είτε Tourmo>0)
                pct = t_val / z_val
                if pct > 1.0:
                    return 1.0  # Cap at 100%
                return round(pct, 2)

            # 1. Tablet_Use_Truck
            z_trucks_agg = (
                df_zonda.groupby(z_vehicle, as_index=False)
                .agg({z_carrier: "first", z_shipping: "first", z_dist: "sum"})
                .rename(
                    columns={
                        z_vehicle: t_veh_col,
                        z_carrier: "Z_Carrier",
                        z_shipping: "Z_Shipping",
                        z_dist: "Zonda_Km",
                    }
                )
            )
            z_trucks_agg["Zonda_Km"] = z_trucks_agg["Zonda_Km"].round(0)

            tablet_truck = pd.merge(
                t_piv1, z_trucks_agg, on=t_veh_col, how="outer"
            )
            tablet_truck["Tourmo_Km"] = (
                tablet_truck["Tourmo_Km"].fillna(0.0).round(0)
            )
            tablet_truck["Zonda_Km"] = (
                tablet_truck["Zonda_Km"].fillna(0.0).round(0)
            )

            tablet_truck[t_group] = tablet_truck[t_group].fillna(
                tablet_truck["Z_Carrier"]
            )
            tablet_truck[t_parent_group] = tablet_truck[t_parent_group].fillna(
                tablet_truck["Z_Shipping"].map(SHIPPING_TO_PARENT)
            )
            tablet_truck[t_parent_group] = tablet_truck[t_parent_group].fillna(
                tablet_truck["Z_Shipping"]
            )
            tablet_truck.drop(columns=["Z_Carrier", "Z_Shipping"], inplace=True)

            tablet_truck["Ποσοστό Χρήσης Tablet"] = tablet_truck.apply(
                lambda r: calc_usage_pct(r["Tourmo_Km"], r["Zonda_Km"]), axis=1
            )

            tablet_truck = tablet_truck[
                [
                    t_parent_group,
                    t_group,
                    t_veh_col,
                    "Tourmo_Km",
                    "Zonda_Km",
                    "Ποσοστό Χρήσης Tablet",
                ]
            ].sort_values(by=[t_parent_group, t_group, t_veh_col])

            # 2. Tablet_Use_Carrier
            z_carriers_agg = (
                df_zonda.groupby(z_sap, as_index=False)
                .agg({z_carrier: "first", z_shipping: "first", z_dist: "sum"})
                .rename(
                    columns={
                        z_sap: t_ext_id,
                        z_carrier: "Z_Carrier",
                        z_shipping: "Z_Shipping",
                        z_dist: "Zonda_Km",
                    }
                )
            )
            z_carriers_agg["Zonda_Km"] = z_carriers_agg["Zonda_Km"].round(0)

            tablet_carrier = pd.merge(
                t_piv2, z_carriers_agg, on=t_ext_id, how="outer"
            )
            tablet_carrier["Tourmo_Km"] = (
                tablet_carrier["Tourmo_Km"].fillna(0.0).round(0)
            )
            tablet_carrier["Zonda_Km"] = (
                tablet_carrier["Zonda_Km"].fillna(0.0).round(0)
            )

            tablet_carrier[t_group] = tablet_carrier[t_group].fillna(
                tablet_carrier["Z_Carrier"]
            )
            tablet_carrier[t_parent_group] = tablet_carrier[
                t_parent_group
            ].fillna(tablet_carrier["Z_Shipping"].map(SHIPPING_TO_PARENT))
            tablet_carrier[t_parent_group] = tablet_carrier[
                t_parent_group
            ].fillna(tablet_carrier["Z_Shipping"])
            tablet_carrier.drop(
                columns=["Z_Carrier", "Z_Shipping"], inplace=True
            )

            tablet_carrier["Ποσοστό Χρήσης Tablet"] = tablet_carrier.apply(
                lambda r: calc_usage_pct(r["Tourmo_Km"], r["Zonda_Km"]), axis=1
            )

            tablet_carrier = tablet_carrier[
                [
                    t_parent_group,
                    t_ext_id,
                    t_group,
                    "Tourmo_Km",
                    "Zonda_Km",
                    "Ποσοστό Χρήσης Tablet",
                ]
            ].sort_values(by=[t_parent_group, t_group])

            # --- E. EXCEL WRITING & FORMATTING ---
            output_buffer = io.BytesIO()

            # 🎨 ΧΡΩΜΑΤΑ ΣΗΜΑΝΣΗΣ (TRAFFIC LIGHT SYSTEM)
            # Κόκκινο (0%)
            red_fill = PatternFill(
                start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
            )
            red_font = Font(color="9C0006", bold=True)

            # Κίτρινο (< 30%)
            yellow_fill = PatternFill(
                start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"
            )
            yellow_font = Font(color="9C6500", bold=True)

            # Πράσινο (≥ 90%)
            green_fill = PatternFill(
                start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
            )
            green_font = Font(color="006100", bold=True)

            center_alignment = Alignment(
                horizontal="center", vertical="center"
            )

            std_header_fill = PatternFill(
                start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"
            )
            std_header_font = Font(
                name="Calibri", size=11, bold=True, color="1F497D"
            )

            tablet_header_fill = PatternFill(
                start_color="BEE3F8", end_color="BEE3F8", fill_type="solid"
            )
            tablet_header_font = Font(
                name="Calibri", size=11, bold=True, color="003E6B"
            )

            with pd.ExcelWriter(output_buffer, engine="openpyxl") as writer:
                # 🌟 ΠΡΩΤΕΣ ΟΙ 2 ΚΑΡΤΕΛΕΣ TABLET USAGE 🌟
                tablet_truck.to_excel(
                    writer, index=False, sheet_name="Tablet_Use_Truck"
                )
                tablet_carrier.to_excel(
                    writer, index=False, sheet_name="Tablet_Use_Carrier"
                )

                # Υπόλοιπες καρτέλες
                df_zonda.to_excel(
                    writer, index=False, sheet_name="Zonda_Orders"
                )
                z_piv1.to_excel(
                    writer, index=False, sheet_name="Zonda_Pivot_Vehicle"
                )
                z_piv2.to_excel(
                    writer, index=False, sheet_name="Zonda_Pivot_SAP"
                )

                df_tourmo.to_excel(
                    writer, index=False, sheet_name="Tourmo_Data"
                )
                t_piv1.to_excel(
                    writer, index=False, sheet_name="Tourmo_Pivot_Vehicle"
                )
                t_piv2.to_excel(
                    writer, index=False, sheet_name="Tourmo_Pivot_SAP"
                )

                for sheet_name in writer.sheets.keys():
                    ws = writer.sheets[sheet_name]
                    ws.freeze_panes = "A2"
                    ws.auto_filter.ref = ws.dimensions

                    is_tablet_tab = sheet_name in [
                        "Tablet_Use_Truck",
                        "Tablet_Use_Carrier",
                    ]
                    h_fill = (
                        tablet_header_fill
                        if is_tablet_tab
                        else std_header_fill
                    )
                    h_font = (
                        tablet_header_font
                        if is_tablet_tab
                        else std_header_font
                    )

                    if is_tablet_tab:
                        ws.sheet_properties.tabColor = "0072CE"

                    for cell in ws[1]:
                        cell.fill = h_fill
                        cell.font = h_font
                        cell.alignment = center_alignment

                    for row in ws.iter_rows(
                        min_row=2,
                        max_row=ws.max_row,
                        min_col=1,
                        max_col=ws.max_column,
                    ):
                        for cell in row:
                            cell.alignment = center_alignment

                    # 🎨 ΧΡΩΜΑΤΙΣΜΟΣ ΠΟΣΟΣΤΩΝ ΣΤΙΣ ΚΑΡΤΕΛΕΣ TABLET USAGE
                    if is_tablet_tab:
                        pct_col_idx = None
                        for col_idx, col in enumerate(ws.columns, start=1):
                            if "ΠΟΣΟΣΤΟ" in str(col[0].value or "").upper():
                                pct_col_idx = col_idx
                                break
                        if not pct_col_idx:
                            pct_col_idx = ws.max_column

                        for row_idx in range(2, ws.max_row + 1):
                            cell = ws.cell(row=row_idx, column=pct_col_idx)
                            cell.number_format = "0%"

                            if cell.value is not None and isinstance(
                                cell.value, (int, float)
                            ):
                                val = cell.value
                                # 🔴 0% -> ΚΟΚΚΙΝΟ
                                if val == 0:
                                    cell.fill = red_fill
                                    cell.font = red_font
                                # 🟡 < 30% -> ΚΙΤΡΙΝΟ
                                elif val < 0.30:
                                    cell.fill = yellow_fill
                                    cell.font = yellow_font
                                # 🟢 ≥ 90% -> ΠΡΑΣΙΝΟ
                                elif val >= 0.90:
                                    cell.fill = green_fill
                                    cell.font = green_font

                    # Πλάτη στηλών
                    for col in ws.columns:
                        h_val = str(col[0].value or "")
                        m_len = max(
                            len(str(cell.value or "")) for cell in col
                        )
                        c_letter = openpyxl.utils.get_column_letter(
                            col[0].column
                        )
                        if (
                            "COMM_CARR" in h_val.upper()
                            or "ΕΞΩΤΕΡΙΚΟ" in h_val.upper()
                            or "ΑΝΑΓΝΩΡΙΣΤΙΚΟ" in h_val.upper()
                        ):
                            ws.column_dimensions[c_letter].width = max(
                                m_len + 8, 28
                            )
                        elif (
                            "TABLET" in h_val.upper()
                            or "ΠΟΣΟΣΤΟ" in h_val.upper()
                        ):
                            ws.column_dimensions[c_letter].width = max(
                                m_len + 6, 24
                            )
                        else:
                            ws.column_dimensions[c_letter].width = max(
                                m_len + 5, 14
                            )

                # Κοκκίνισμα Zonda Orders για αλλαγές αποστάσεων
                ws_z = writer.sheets["Zonda_Orders"]
                d_idx = None
                for idx, col_name in enumerate(df_zonda.columns, start=1):
                    if col_name == z_dist:
                        d_idx = idx
                        break
                if d_idx:
                    for row_idx in range(2, ws_z.max_row + 1):
                        if (row_idx - 2) < len(is_dist_changed_zonda):
                            if is_dist_changed_zonda[row_idx - 2]:
                                d_cell = ws_z.cell(
                                    row=row_idx, column=d_idx
                                )
                                d_cell.fill = red_fill
                                d_cell.font = red_font

            st.success("🎉 **Η επεξεργασία ολοκληρώθηκε επιτυχώς!**")

            # KPIs
            k1, k2, k3, k4 = st.columns(4)
            with k1:
                st.metric("Παραγγελίες Zonda", len(df_zonda))
            with k2:
                st.metric("Σύνολο Οχημάτων (Έλεγχος)", len(tablet_truck))
            with k3:
                st.metric(
                    "Χιλιόμετρα Zonda (x2)",
                    f"{int(round(df_zonda[z_dist].sum())):,}",
                )
            with k4:
                avg_u = tablet_truck["Ποσοστό Χρήσης Tablet"].mean() * 100
                st.metric("Μέση Χρήση Tablet", f"{avg_u:.0f}%")

            st.markdown("<br>", unsafe_allow_html=True)

            # Download Button (v4.1)
            st.download_button(
                label="📥 Λήψη Αναφοράς Excel (Logistics_Report_v4_COMPLETE.xlsx)",
                data=output_buffer.getvalue(),
                file_name="Logistics_Report_v4_COMPLETE.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
